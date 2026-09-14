import os
import openpyxl

def create_single_sheet_multi_table(filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    # Table 1: Customers
    ws.cell(row=1, column=1, value="Customers")
    headers1 = ["id", "name", "city"]
    for col_idx, h in enumerate(headers1, 1):
        ws.cell(row=2, column=col_idx, value=h)
    rows1 = [
        ["1", "Alice", "Rome"],
        ["2", "Bob", "Milan"],
        ["3", "Charlie", "Turin"],
    ]
    for r_idx, row in enumerate(rows1, 3):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Table 2: Orders
    ws.cell(row=8, column=1, value="Orders")
    headers2 = ["order_id", "customer_id", "amount", "order_date"]
    for col_idx, h in enumerate(headers2, 1):
        ws.cell(row=9, column=col_idx, value=h)
    rows2 = [
        ["101", "1", "250.50", "2023-01-15"],
        ["102", "2", "89.00", "2023-01-16"],
        ["103", "1", "1200.00", "2023-01-20"],
    ]
    for r_idx, row in enumerate(rows2, 10):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)
    print(f"Created {filepath}")

def create_multi_sheet_multi_table(filepath):
    wb = openpyxl.Workbook()
    
    # Sheet 1: Sales
    ws_sales = wb.active
    ws_sales.title = "Sales"

    # Table 1.1: Orders
    ws_sales.cell(row=1, column=1, value="Orders")
    headers_orders = ["order_id", "customer_id", "product_id", "amount"]
    for col_idx, h in enumerate(headers_orders, 1):
        ws_sales.cell(row=2, column=col_idx, value=h)
    rows_orders = [
        ["101", "1", "1", "250.50"],
        ["102", "2", "2", "140.00"],
        ["103", "1", "1", "85.20"],
    ]
    for r_idx, row in enumerate(rows_orders, 3):
        for c_idx, val in enumerate(row, 1):
            ws_sales.cell(row=r_idx, column=c_idx, value=val)

    # Table 1.2: Returns
    ws_sales.cell(row=9, column=1, value="Returns")
    headers_returns = ["return_id", "order_id", "reason"]
    for col_idx, h in enumerate(headers_returns, 1):
        ws_sales.cell(row=10, column=col_idx, value=h)
    rows_returns = [
        ["501", "102", "Defective"],
        ["502", "101", "Wrong size"],
        ["503", "103", "Unwanted"],
    ]
    for r_idx, row in enumerate(rows_returns, 11):
        for c_idx, val in enumerate(row, 1):
            ws_sales.cell(row=r_idx, column=c_idx, value=val)

    # Sheet 2: Inventory
    ws_inv = wb.create_sheet(title="Inventory")

    # Table 2.1: Products
    ws_inv.cell(row=1, column=1, value="Products")
    headers_products = ["product_id", "supplier_id", "name", "price", "stock"]
    for col_idx, h in enumerate(headers_products, 1):
        ws_inv.cell(row=2, column=col_idx, value=h)
    rows_products = [
        ["1", "1", "Laptop", "999.99", "15"],
        ["2", "2", "Mouse", "25.50", "120"],
        ["3", "1", "Keyboard", "45.00", "80"],
    ]
    for r_idx, row in enumerate(rows_products, 3):
        for c_idx, val in enumerate(row, 1):
            ws_inv.cell(row=r_idx, column=c_idx, value=val)

    # Table 2.2: Suppliers
    ws_inv.cell(row=9, column=1, value="Suppliers")
    headers_suppliers = ["supplier_id", "company", "country"]
    for col_idx, h in enumerate(headers_suppliers, 1):
        ws_inv.cell(row=10, column=col_idx, value=h)
    rows_suppliers = [
        ["1", "TechCorp", "USA"],
        ["2", "EuroParts", "Germany"],
        ["3", "AsiaSupply", "Japan"],
    ]
    for r_idx, row in enumerate(rows_suppliers, 11):
        for c_idx, val in enumerate(row, 1):
            ws_inv.cell(row=r_idx, column=c_idx, value=val)

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)
    print(f"Created {filepath}")

if __name__ == "__main__":
    fixtures_dir = os.path.join(os.path.dirname(__file__), "..", "test", "fixtures", "excel")
    root_fixtures_dir = os.path.join(os.path.dirname(__file__), "..", "..", "test_fixtures")

    create_single_sheet_multi_table(os.path.join(fixtures_dir, "multi_table_same_sheet.xlsx"))
    create_single_sheet_multi_table(os.path.join(fixtures_dir, "single_sheet_multi_table.xlsx"))
    create_single_sheet_multi_table(os.path.join(root_fixtures_dir, "single_sheet_multi_table.xlsx"))

    create_multi_sheet_multi_table(os.path.join(fixtures_dir, "multi_sheet_multi_table.xlsx"))
    create_multi_sheet_multi_table(os.path.join(root_fixtures_dir, "multi_sheet_multi_table.xlsx"))
